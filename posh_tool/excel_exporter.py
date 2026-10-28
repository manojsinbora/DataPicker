"""
Excel Exporter using openpyxl.
Strictly implements:
- Sheet 1: "Companies" with exact 21 columns
- Sheet 2: "Notes & Sources"
- Bold, frozen header row
- Auto-filters on all columns
- Sensible column widths & text wrapping
- Clickable hyperlinks for URLs
- Continuous row numbering (#)
- Appending across batches into the same single workbook
"""

import os
from typing import List, Dict, Any, Optional
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .config import COLUMNS, NOT_AVAILABLE


class ExcelExporter:
    def __init__(self, city: str, start_date: Optional[str] = None):
        self.city = city.strip().title()
        self.date_str = start_date or datetime.now().strftime("%Y-%m-%d")
        # Format: POSH_IC_Contacts_<City>_<YYYY-MM-DD>.xlsx
        clean_city = self.city.replace(" ", "_")
        self.filename = f"POSH_IC_Contacts_{clean_city}_{self.date_str}.xlsx"

    def export_batch(
        self,
        new_companies: List[Dict[str, Any]],
        area: str,
        industry: str,
        target_rows: int,
        batch_no: int,
        excluded_companies: List[Dict[str, str]],
        all_source_urls: List[str],
        require_phone: str = "optional"
    ) -> str:
        """
        Creates or appends to the Excel file, saving all batches into the SAME file.
        Returns the absolute filepath of the Excel workbook.
        """
        # Load existing workbook or create new
        if os.path.exists(self.filename):
            wb = openpyxl.load_workbook(self.filename)
            ws_companies = wb["Companies"]
            ws_notes = wb["Notes & Sources"]
            start_row = ws_companies.max_row + 1
            existing_count = ws_companies.max_row - 1  # exclude header
        else:
            wb = openpyxl.Workbook()
            ws_companies = wb.active
            ws_companies.title = "Companies"
            ws_notes = wb.create_sheet(title="Notes & Sources")
            start_row = 2
            existing_count = 0
            # Setup Companies header
            self._setup_companies_header(ws_companies)

        # Style definitions
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        regular_alignment = Alignment(vertical="top", wrap_text=True)

        current_num = existing_count + 1

        for comp in new_companies:
            row_data = [
                current_num,
                comp.get("Company name") or NOT_AVAILABLE,
                comp.get("Legal form (Pvt Ltd / LLP / Listed / Not confirmed)") or NOT_AVAILABLE,
                comp.get("Full address with PIN") or NOT_AVAILABLE,
                comp.get("Company website") or NOT_AVAILABLE,
                comp.get("Company public phone number") or NOT_AVAILABLE,
                comp.get("Company public email ID") or NOT_AVAILABLE,
                comp.get("POSH / IC web page link") or NOT_AVAILABLE,
                comp.get("IC document / report link (POSH policy PDF, IC notice, annual report or BRSR where IC details appear)") or NOT_AVAILABLE,
                comp.get("Where IC details were found (document name + page/section)") or NOT_AVAILABLE,
                comp.get("IC presiding officer / member name(s) (as published)") or NOT_AVAILABLE,
                comp.get("IC email ID") or NOT_AVAILABLE,
                comp.get("IC phone / mobile number") or NOT_AVAILABLE,
                comp.get("HR contact name (as published, if any)") or NOT_AVAILABLE,
                comp.get("HR email ID") or NOT_AVAILABLE,
                comp.get("HR phone number") or NOT_AVAILABLE,
                comp.get("Employee count evidence") or NOT_AVAILABLE,
                comp.get("Source URL(s)") or NOT_AVAILABLE,
                comp.get("Date of source / date checked") or datetime.now().strftime("%Y-%m-%d"),
                batch_no,
                comp.get("Notes") or NOT_AVAILABLE
            ]

            # Write row cells
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws_companies.cell(row=start_row, column=col_idx)
                
                # Integer columns: # and Batch no.
                if col_idx in (1, 20):
                    try:
                        cell.value = int(val)
                    except (ValueError, TypeError):
                        cell.value = val
                    cell.font = regular_font
                    cell.alignment = regular_alignment
                    continue

                str_val = str(val).strip() if val is not None else NOT_AVAILABLE
                if not str_val:
                    str_val = NOT_AVAILABLE
                cell.value = str_val
                cell.font = regular_font
                cell.alignment = regular_alignment

                # Turn URLs into clickable links
                if str_val.startswith("http://") or str_val.startswith("https://"):
                    cell.hyperlink = str_val
                    cell.font = link_font

            start_row += 1
            current_num += 1

        # Freeze header row & apply auto filter
        ws_companies.freeze_panes = "A2"
        max_col_letter = get_column_letter(len(COLUMNS))
        ws_companies.auto_filter.ref = f"A1:{max_col_letter}{ws_companies.max_row}"

        # Adjust column widths
        self._adjust_column_widths(ws_companies)

        # Update Sheet 2: "Notes & Sources"
        total_delivered = ws_companies.max_row - 1
        self._update_notes_sheet(
            ws_notes=ws_notes,
            ws_companies=ws_companies,
            city=self.city,
            area=area,
            industry=industry,
            target_rows=target_rows,
            total_delivered=total_delivered,
            start_date=self.date_str,
            excluded_companies=excluded_companies,
            all_source_urls=all_source_urls,
            require_phone=require_phone
        )

        # Save workbook
        wb.save(self.filename)
        return os.path.abspath(self.filename)

    def _setup_companies_header(self, ws):
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        ws.row_dimensions[1].height = 35

    def _adjust_column_widths(self, ws):
        widths = {
            1: 6,   # #
            2: 26,  # Company name
            3: 16,  # Legal form
            4: 35,  # Address
            5: 28,  # Website
            6: 20,  # Phone
            7: 26,  # Email
            8: 30,  # POSH page
            9: 35,  # IC doc link
            10: 30, # Where found
            11: 30, # Members
            12: 26, # IC email
            13: 20, # IC phone
            14: 24, # HR name
            15: 26, # HR email
            16: 20, # HR phone
            17: 30, # Employee evidence
            18: 35, # Source URLs
            19: 16, # Date
            20: 10, # Batch no.
            21: 25  # Notes
        }
        for col_idx, width in widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

    def _update_notes_sheet(
        self,
        ws_notes,
        ws_companies,
        city: str,
        area: str,
        industry: str,
        target_rows: int,
        total_delivered: int,
        start_date: str,
        excluded_companies: List[Dict[str, str]],
        all_source_urls: List[str],
        require_phone: str = "optional"
    ):
        """
        Populates Sheet 2 "Notes & Sources" according to exact specifications.
        """
        # Clear existing contents
        for row in ws_notes.iter_rows(max_col=10, max_row=ws_notes.max_row):
            for cell in row:
                cell.value = None

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        section_font = Font(name="Calibri", size=12, bold=True, color="1F497D")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        meta_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

        # 1. Project Metadata
        ws_notes.cell(row=1, column=1, value="PROJECT METADATA").font = section_font
        meta = [
            ("City", city),
            ("Area / Locality", area),
            ("Industry Filter", industry),
            ("Phone Requirement", require_phone.title() if require_phone else "Optional"),
            ("Row Target", target_rows),
            ("Rows Delivered So Far", total_delivered),
            ("Research Start Date", start_date),
            ("Last Updated", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        ]
        for idx, (label, val) in enumerate(meta, start=2):
            c1 = ws_notes.cell(row=idx, column=1, value=label)
            c2 = ws_notes.cell(row=idx, column=2, value=val)
            c1.font = bold_font
            c1.fill = meta_fill
            c2.font = regular_font

        # 2. Running Gaps Summary
        current_row = len(meta) + 4
        ws_notes.cell(row=current_row, column=1, value="RUNNING GAPS SUMMARY").font = section_font
        current_row += 1

        # Calculate counts from ws_companies
        total_rows = ws_companies.max_row - 1
        gap_fields = {
            "IC email ID": 12,
            "IC phone / mobile number": 13,
            "HR email ID": 15,
            "HR phone number": 16,
            "Company public phone number": 6,
            "Company public email ID": 7,
            "IC document / report link": 9
        }

        # Header for gaps table
        ws_notes.cell(row=current_row, column=1, value="Field Name").font = header_font
        ws_notes.cell(row=current_row, column=1).fill = header_fill
        ws_notes.cell(row=current_row, column=2, value="Found Count").font = header_font
        ws_notes.cell(row=current_row, column=2).fill = header_fill
        ws_notes.cell(row=current_row, column=3, value="Total Checked").font = header_font
        ws_notes.cell(row=current_row, column=3).fill = header_fill
        ws_notes.cell(row=current_row, column=4, value="Completion Rate").font = header_font
        ws_notes.cell(row=current_row, column=4).fill = header_fill
        current_row += 1

        for field_name, col_idx in gap_fields.items():
            found_count = 0
            if total_rows > 0:
                for r in range(2, ws_companies.max_row + 1):
                    v = str(ws_companies.cell(row=r, column=col_idx).value or "").strip()
                    if v and v != NOT_AVAILABLE:
                        found_count += 1
            pct = f"{(found_count / total_rows * 100):.1f}%" if total_rows > 0 else "0%"

            ws_notes.cell(row=current_row, column=1, value=field_name).font = regular_font
            ws_notes.cell(row=current_row, column=2, value=found_count).font = regular_font
            ws_notes.cell(row=current_row, column=3, value=total_rows).font = regular_font
            ws_notes.cell(row=current_row, column=4, value=pct).font = regular_font
            current_row += 1

        # 3. Companies Checked and Excluded (with one-line reasons)
        current_row += 2
        ws_notes.cell(row=current_row, column=1, value="COMPANIES CHECKED AND EXCLUDED").font = section_font
        current_row += 1

        ws_notes.cell(row=current_row, column=1, value="Company / Entity Name").font = header_font
        ws_notes.cell(row=current_row, column=1).fill = header_fill
        ws_notes.cell(row=current_row, column=2, value="Exclusion Reason").font = header_font
        ws_notes.cell(row=current_row, column=2).fill = header_fill
        current_row += 1

        if excluded_companies:
            for exc in excluded_companies:
                ws_notes.cell(row=current_row, column=1, value=exc.get("name", "Unknown")).font = regular_font
                ws_notes.cell(row=current_row, column=2, value=exc.get("reason", "Did not meet criteria")).font = regular_font
                current_row += 1
        else:
            ws_notes.cell(row=current_row, column=1, value="None").font = regular_font
            ws_notes.cell(row=current_row, column=2, value="All checked companies met qualification rules").font = regular_font
            current_row += 1

        # 4. Full List of Source URLs
        current_row += 2
        ws_notes.cell(row=current_row, column=1, value="FULL LIST OF SOURCE URLS").font = section_font
        current_row += 1

        ws_notes.cell(row=current_row, column=1, value="#").font = header_font
        ws_notes.cell(row=current_row, column=1).fill = header_fill
        ws_notes.cell(row=current_row, column=2, value="Source URL").font = header_font
        ws_notes.cell(row=current_row, column=2).fill = header_fill
        current_row += 1

        unique_sources = sorted(list(set(all_source_urls)))
        for s_idx, s_url in enumerate(unique_sources, start=1):
            ws_notes.cell(row=current_row, column=1, value=s_idx).font = regular_font
            c_url = ws_notes.cell(row=current_row, column=2, value=s_url)
            c_url.font = link_font
            c_url.hyperlink = s_url
            current_row += 1

        # Adjust column widths on Notes sheet
        ws_notes.column_dimensions["A"].width = 32
        ws_notes.column_dimensions["B"].width = 50
        ws_notes.column_dimensions["C"].width = 16
        ws_notes.column_dimensions["D"].width = 16

